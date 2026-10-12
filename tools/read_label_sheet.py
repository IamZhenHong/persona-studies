"""Read the filled labelling spreadsheet into results/validation/hand_labels.json (what validate_judge.report reads).
Rows count only when all four main traits are filled; Justified refusal blank = No. Checks the refusal/withholding rule.
  uv run --with openpyxl python persona/tools/read_label_sheet.py [path.xlsx]"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from persona.src.common import PERSONA  # noqa: E402

V = PERSONA / "results" / "validation"
path = Path(sys.argv[1]) if len(sys.argv) > 1 else V / "hand_label_sheet.xlsx"
ws = load_workbook(path, data_only=True)["Labels"]
head = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(head)}
keys = {"Warmth": "warmth", "Helpful": "helpfulness", "Contempt": "contempt", "Withholding": "withholding",
        "Justified refusal": "justified_refusal"}
yn = {"yes": 1, "no": 0, "y": 1, "n": 0, "1": 1, "0": 0}
out, partial, bad, breaks = {}, 0, [], []
for row in ws.iter_rows(min_row=2, values_only=True):
    iid = row[col["Item ID"]]
    vals = {}
    for h, k in keys.items():
        v = row[col[h]]
        if v is None or str(v).strip() == "":
            continue
        s = str(v).strip().lower()
        if s not in yn:
            bad.append((iid, h, v)); continue
        vals[k] = yn[s]
    main = ("warmth", "helpfulness", "contempt", "withholding")
    if all(k in vals for k in main):
        vals.setdefault("justified_refusal", 0)
        if vals["justified_refusal"] and vals["withholding"]:
            breaks.append(iid)
        out[iid] = vals
    elif vals:
        partial += 1
(V / "hand_labels.json").write_text(json.dumps(out, indent=1))
print(f"{len(out)} fully labelled rows -> {V / 'hand_labels.json'}; {partial} partly filled (ignored); "
      f"{len(bad)} unreadable cells {bad[:5]}; {len(breaks)} rule breaks (refusal and withholding both Yes) {breaks[:5]}")
